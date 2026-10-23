import re
from io import BytesIO
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="CleanReport Studio", page_icon="🧹", layout="wide")
st.markdown("""
<style>
.stApp{background:#f6f8fc}.block-container{padding-top:1.5rem;max-width:1450px}
.hero{padding:1.4rem 1.6rem;border-radius:18px;background:linear-gradient(120deg,#142d4e,#2563eb);color:white;margin-bottom:1rem}
.hero h1{color:white;margin:0}.hero p{color:#e4edff;margin:.45rem 0 0}
div[data-testid="stMetric"]{background:white;border:1px solid #e3e9f3;border-radius:12px;padding:12px}
</style>
<div class="hero"><h1>🧹 CleanReport Studio</h1>
<p>Clean messy spreadsheets, inspect data quality, and generate an automated Excel report.</p></div>
""", unsafe_allow_html=True)

def clean_col_name(value):
    value = re.sub(r"[^a-z0-9]+", "_", str(value).strip().lower())
    value = re.sub(r"_+", "_", value).strip("_")
    return value or "unnamed_column"

def load_file(upload):
    if Path(upload.name).suffix.lower() == ".csv":
        return pd.read_csv(upload)
    return pd.read_excel(upload)

def profile(df):
    result = []
    for col in df.columns:
        s = df[col]
        result.append({
            "Column": str(col), "Data Type": str(s.dtype),
            "Missing Values": int(s.isna().sum()),
            "Missing %": round(float(s.isna().mean() * 100), 2) if len(s) else 0,
            "Unique Values": int(s.nunique(dropna=True)),
            "Duplicate Values": int(s.duplicated().sum())
        })
    return pd.DataFrame(result)

def clean_data(source, normalize_cols, trim_text, text_case, missing_strategy, constant, remove_duplicates):
    df = source.copy()
    stats = {"rows_before": len(df), "cols_before": len(df.columns),
             "missing_before": int(df.isna().sum().sum()), "dupes_before": int(df.duplicated().sum())}
    actions = []
    if normalize_cols:
        new_names, used = [], {}
        for col in df.columns:
            name = clean_col_name(col)
            used[name] = used.get(name, 0) + 1
            new_names.append(name if used[name] == 1 else f"{name}_{used[name]}")
        df.columns = new_names
        actions.append("Standardized column names")
    for col in df.select_dtypes(include=["object", "string"]).columns:
        if trim_text:
            df[col] = df[col].astype("string").str.strip().replace("", pd.NA)
        if text_case == "lowercase":
            df[col] = df[col].astype("string").str.lower()
        elif text_case == "Title Case":
            df[col] = df[col].astype("string").str.title()
    if trim_text:
        actions.append("Trimmed extra spaces and converted empty text to missing")
    if missing_strategy == "Fill with median/mode":
        for col in df.columns:
            if df[col].isna().any():
                if pd.api.types.is_numeric_dtype(df[col]):
                    df[col] = df[col].fillna(df[col].median())
                else:
                    modes = df[col].mode(dropna=True)
                    df[col] = df[col].fillna(modes.iloc[0] if not modes.empty else "Unknown")
        actions.append("Filled numeric blanks with median and text blanks with mode/Unknown")
    elif missing_strategy == "Fill with a constant":
        for col in df.columns:
            if df[col].isna().any():
                fill = constant
                if pd.api.types.is_numeric_dtype(df[col]):
                    try: fill = float(constant)
                    except ValueError: fill = 0
                df[col] = df[col].fillna(fill)
        actions.append(f"Filled missing values with {constant!r} (numeric blanks use its numeric value or 0)")
    elif missing_strategy == "Drop rows with missing values":
        before = len(df)
        df = df.dropna()
        actions.append(f"Dropped {before - len(df)} rows with missing values")
    if remove_duplicates:
        before = len(df)
        df = df.drop_duplicates().reset_index(drop=True)
        actions.append(f"Removed {before - len(df)} duplicate rows")
    stats.update({"rows_after": len(df), "cols_after": len(df.columns),
                  "missing_after": int(df.isna().sum().sum()), "dupes_after": int(df.duplicated().sum()),
                  "actions": actions})
    return df, stats

def excel_report(cleaned, original, stats):
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        cleaned.to_excel(writer, sheet_name="Cleaned Data", index=False)
        profile(original).to_excel(writer, sheet_name="Data Quality", index=False)
        pd.DataFrame([
            ["Rows before", stats["rows_before"]], ["Rows after", stats["rows_after"]],
            ["Columns before", stats["cols_before"]], ["Columns after", stats["cols_after"]],
            ["Missing cells before", stats["missing_before"]], ["Missing cells after", stats["missing_after"]],
            ["Duplicate rows before", stats["dupes_before"]], ["Duplicate rows after", stats["dupes_after"]]
        ], columns=["Metric", "Value"]).to_excel(writer, sheet_name="Summary", index=False)
        pd.DataFrame({"Cleaning Actions": stats["actions"] or ["No cleaning actions selected"]}).to_excel(writer, sheet_name="Actions", index=False)
        wb = writer.book
        fill = PatternFill("solid", fgColor="1D4ED8")
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = fill
                cell.alignment = Alignment(vertical="center")
            for cells in ws.columns:
                width = min(max(max((len(str(c.value)) if c.value is not None else 0) for c in cells) + 2, 12), 42)
                ws.column_dimensions[get_column_letter(cells[0].column)].width = width
        ws = wb["Summary"]
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Before vs After"
        chart.style = 10
        chart.add_data(Reference(ws, min_col=2, min_row=2, max_row=5), titles_from_data=False)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
        chart.height, chart.width = 6, 12
        ws.add_chart(chart, "D2")
    return buf.getvalue()

with st.sidebar:
    st.header("1 · Upload data")
    uploaded = st.file_uploader("Choose CSV or Excel", type=["csv", "xlsx", "xls"])
    st.caption("Use sample file: data/sample_sales.csv")
    st.divider()
    st.header("2 · Cleaning rules")
    normalize_cols = st.checkbox("Standardize column names", True)
    trim_text = st.checkbox("Trim text / empty strings", True)
    remove_duplicates = st.checkbox("Remove duplicate rows", True)
    text_case = st.selectbox("Text case", ["Keep as-is", "lowercase", "Title Case"])
    missing_strategy = st.selectbox("Missing values", ["Keep missing values", "Fill with median/mode", "Fill with a constant", "Drop rows with missing values"])
    constant = st.text_input("Constant fill value", "Unknown")
    run = st.button("✨ Clean data & build report", type="primary", use_container_width=True)

if uploaded is None:
    left, right = st.columns([1.2, 1])
    with left:
        st.subheader("From messy spreadsheet to clean report")
        st.write("Upload CSV/Excel data, inspect quality issues, apply chosen cleaning rules, and export a formatted workbook.")
        st.markdown("- Data quality profile\n- Missing-value and duplicate handling\n- Before/after comparison\n- Numeric visualizations\n- CSV and multi-sheet Excel downloads")
    with right:
        st.info("Try the included `data/sample_sales.csv` file.")
        st.code("python -m pip install -r requirements.txt\npython -m streamlit run app.py", language="powershell")
    st.stop()

try:
    original = load_file(uploaded)
except Exception as exc:
    st.error(f"Could not read this file: {exc}")
    st.stop()
if original.empty:
    st.warning("This file has no data rows.")
    st.stop()

st.subheader("Raw data preview")
st.caption(f"{uploaded.name} · {len(original):,} rows · {len(original.columns):,} columns")
st.dataframe(original.head(100), use_container_width=True, hide_index=True)
quality = profile(original)
a,b,c,d = st.columns(4)
a.metric("Rows", f"{len(original):,}")
b.metric("Columns", f"{len(original.columns):,}")
c.metric("Missing cells", f"{int(original.isna().sum().sum()):,}")
d.metric("Duplicate rows", f"{int(original.duplicated().sum()):,}")
with st.expander("Data quality by column", expanded=True):
    st.dataframe(quality, use_container_width=True, hide_index=True)

if run:
    cleaned, stats = clean_data(original, normalize_cols, trim_text, text_case, missing_strategy, constant, remove_duplicates)
    st.session_state["cleaned"] = cleaned
    st.session_state["stats"] = stats
    st.session_state["excel"] = excel_report(cleaned, original, stats)

if "cleaned" in st.session_state:
    cleaned, stats = st.session_state["cleaned"], st.session_state["stats"]
    st.divider()
    st.subheader("Cleaning results")
    a,b,c,d = st.columns(4)
    a.metric("Rows after", f"{stats['rows_after']:,}", delta=f"{stats['rows_after']-stats['rows_before']:,} rows")
    b.metric("Missing after", f"{stats['missing_after']:,}", delta=f"{stats['missing_after']-stats['missing_before']:,}")
    c.metric("Duplicates after", f"{stats['dupes_after']:,}", delta=f"{stats['dupes_after']-stats['dupes_before']:,}")
    d.metric("Rows removed", f"{stats['rows_before']-stats['rows_after']:,}")
    st.write("**Actions applied**")
    for action in stats["actions"] or ["No cleaning actions selected"]:
        st.write("✓ " + action)
    t1,t2,t3 = st.tabs(["Cleaned data", "Visual insights", "Download reports"])
    with t1:
        st.dataframe(cleaned.head(300), use_container_width=True, hide_index=True)
    with t2:
        numeric = cleaned.select_dtypes(include=np.number).columns.tolist()
        if numeric:
            selected = st.selectbox("Numeric column to visualize", numeric)
            fig, ax = plt.subplots(figsize=(9,4))
            cleaned[selected].dropna().plot(kind="hist", bins=20, ax=ax, edgecolor="white")
            ax.set_title(f"Distribution of {selected}")
            ax.set_xlabel(selected)
            ax.set_ylabel("Frequency")
            st.pyplot(fig)
            plt.close(fig)
        else:
            st.info("No numeric columns are available for a histogram.")
        missing = cleaned.isna().sum().sort_values(ascending=False)
        if missing.sum():
            st.bar_chart(missing[missing > 0])
        else:
            st.success("No missing cells remain.")
    with t3:
        st.download_button("⬇ Download cleaned CSV", cleaned.to_csv(index=False).encode("utf-8-sig"), "cleaned_data.csv", "text/csv")
        st.download_button("⬇ Download Excel report", st.session_state["excel"], "automated_cleaning_report.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
        st.caption("Excel includes Cleaned Data, Data Quality, Summary, and Actions sheets.")
else:
    st.info("Select your cleaning options in the sidebar, then click **Clean data & build report**.")
